"""
로그에서 관측 가능한 이벤트 목록을 뽑는다.

    python -m harness.catalog "C:/qa_logs/CRX3212"
    python -m harness.catalog <로그> --tag LogWriter

시트의 VALUE 를 채울 때 "이 제품에서 실제로 뭐가 찍히는가"를 먼저 본다.
Jira 이슈 제목만 보고 문자열을 지어내면 영원히 미검출로 실패한다.

이미 시트에 쓰이고 있는 문자열은 [사용중] 으로 표시하므로,
아직 안 쓰는 이벤트가 자동화 후보로 남는다.
"""

from __future__ import annotations

import argparse
import collections
import re
from pathlib import Path

ROOT = Path(__file__).parent.parent

# 하네스 형식과 기존 캡처 형식 모두에서 본문만 뽑는다
BODY = [re.compile(r"\|\s*[\w.-]+\s*\|\s?(.*)$"),
        re.compile(r"^\[[^\]]+\]\s?(.*)$")]
EVENT = re.compile(r"\(([A-Za-z][\w]{2,14})\)\s*([^\r\n]{0,80})")


def body_of(line: str) -> str:
    for rx in BODY:
        m = rx.search(line)
        if m:
            return m.group(1)
    return line


def norm(text: str) -> str:
    t = re.sub(r"\b\d[\d.:]*\b", "N", text)
    return re.sub(r"\s+", " ", t).strip()[:62]


def used_values(spec: Path) -> set[str]:
    try:
        import openpyxl
    except ImportError:
        return set()
    if not spec.exists():
        return set()
    wb = openpyxl.load_workbook(spec, data_only=True, read_only=True)
    ws = wb["SPEC"] if "SPEC" in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h else "" for h in next(it)]
    if "VALUE" not in hdr:
        return set()
    vi = hdr.index("VALUE")
    return {str(r[vi]).strip() for r in it if r and r[vi]}


def resolve_log(p: Path) -> Path:
    if p.is_dir():
        mk = p / "latest.txt"
        if mk.exists():
            q = p / mk.read_text(encoding="utf-8").strip()
            if q.exists():
                return q
        subs = [d for d in p.iterdir() if d.is_dir() and any(d.glob("*.log"))]
        cands = list(p.glob("*.log"))
        for d in subs:
            cands += list(d.glob("*.log"))
        if cands:
            return max(cands, key=lambda x: x.stat().st_mtime)
    return p


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("logfile")
    ap.add_argument("--tag", default=None, help="특정 태그만 (예: LogWriter)")
    ap.add_argument("--spec", default=str(ROOT / "spec_v2.xlsx"))
    ap.add_argument("--min", type=int, default=1, help="이 횟수 미만은 숨김")
    a = ap.parse_args()

    p = resolve_log(Path(a.logfile))
    if not p.is_file():
        raise SystemExit(f"로그를 찾을 수 없습니다: {a.logfile}")
    print(f"{p.name}\n")

    used = used_values(Path(a.spec))
    counts: dict[str, collections.Counter] = collections.defaultdict(
        collections.Counter)
    for line in p.open(encoding="utf-8", errors="replace"):
        for m in EVENT.finditer(body_of(line)):
            tag, rest = m.group(1), norm(m.group(2))
            if a.tag and tag != a.tag:
                continue
            counts[tag][f"({tag}) {rest}"] += 1

    # LogWriter 를 먼저. 앱이 남기는 상태 이벤트라 자동화 가치가 가장 높다.
    order = sorted(counts, key=lambda t: (t != "LogWriter", -sum(counts[t].values())))
    total_new = 0
    for tag in order:
        items = [(k, v) for k, v in sorted(counts[tag].items()) if v >= a.min]
        if not items:
            continue
        print(f"=== ({tag})  {len(items)}종 ===")
        for k, v in items:
            mark = "[사용중]" if any(u in k for u in used) else "        "
            if mark.strip() == "":
                total_new += 1
            print(f"  {v:6}  {mark} {k}")
        print()

    print("-" * 66)
    print("[사용중] 은 이미 시트 VALUE 에 들어 있는 문자열입니다.")
    print("나머지가 자동화 후보입니다. 숫자는 N 으로 정규화되어 있으니")
    print("시트에 넣을 때는 변하지 않는 앞부분만 쓰세요.")
    print("  예) (LogWriter) FW_VERSION   <- 버전 값은 빼고")


if __name__ == "__main__":
    main()
