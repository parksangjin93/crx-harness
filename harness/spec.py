"""
테스트 명세 시트 로더.

시트가 진실의 소스다. 코드를 생성하지 않고 런타임에 읽는다.
시트를 고치면 다음 실행에 바로 반영된다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

HEADERS = [
    "출처", "TC_ID", "PRODUCT", "TC_NAME", "TEST_PURPOSE", "STEP",
    "TEST_METHOD", "TYPE", "VALUE", "FAIL", "TIMEOUT", "TIMEOUT_FROM",
    "MATCH", "ORDERED", "SCOPE", "TAG", "CAPTURE", "BLOCKED_BY", "비고",
]


@dataclass
class Step:
    no: int
    type: str            # CHECK | ABSENT | ACTION | SHELL_CHECK
    value: str
    method: str = ""
    fail: str | None = None
    timeout: float | None = None
    timeout_from: str = "ACTION"     # ACTION | PREV_STEP
    match: str = "CONTAINS"
    scope: str = ""                  # ABSENT 전용: TC
    capture: str | None = None       # "이름" 또는 "이름:정규식"


@dataclass
class TestCase:
    tc_id: str
    name: str
    purpose: str
    product: str = "*"
    ordered: bool = True
    tags: list[str] = field(default_factory=list)
    blocked_by: str | None = None
    steps: list[Step] = field(default_factory=list)

    @property
    def actions(self) -> list[Step]:
        return [s for s in self.steps if s.type == "ACTION"]

    @property
    def checks(self) -> list[Step]:
        return [s for s in self.steps if s.type in ("CHECK", "SHELL_CHECK")]

    @property
    def absents(self) -> list[Step]:
        return [s for s in self.steps if s.type == "ABSENT"]

    def applies_to(self, product: str) -> bool:
        return self.product in ("*", product)


class SpecError(Exception):
    """시트 자체가 잘못된 경우. 테스트 실패가 아니라 설정 오류다."""


def load(path: str | Path, sheet: str = "SPEC") -> list[TestCase]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise SpecError(f"시트 '{sheet}' 없음. 있는 시트: {wb.sheetnames}")
    ws = wb[sheet]

    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(it)]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        raise SpecError(f"필수 컬럼 없음: {missing}")
    idx = {h: header.index(h) for h in HEADERS}

    cases: dict[str, TestCase] = {}
    for raw in it:
        if raw is None or not any(v is not None for v in raw):
            continue

        def g(col: str):
            v = raw[idx[col]]
            return v.strip() if isinstance(v, str) else v

        tc_id = str(g("TC_ID") or "").strip()
        if not tc_id:
            continue

        tc = cases.get(tc_id)
        if tc is None:
            tc = TestCase(
                tc_id=tc_id,
                name=str(g("TC_NAME") or tc_id),
                purpose=str(g("TEST_PURPOSE") or ""),
                product=str(g("PRODUCT") or "*"),
                ordered=str(g("ORDERED") or "Y").upper() != "N",
                tags=[t.strip() for t in str(g("TAG") or "").split(",") if t.strip()],
                blocked_by=g("BLOCKED_BY") or None,
            )
            cases[tc_id] = tc
        # 한 TC 안에서 BLOCKED_BY 가 한 행이라도 있으면 전체를 blocked 로 본다
        if g("BLOCKED_BY") and not tc.blocked_by:
            tc.blocked_by = str(g("BLOCKED_BY"))

        typ = str(g("TYPE") or "").upper()
        val = g("VALUE")
        if typ not in ("CHECK", "ABSENT", "ACTION", "SHELL_CHECK") or val is None:
            continue    # 검증 조건이 없는 행 (Jira 원문 등)

        to = g("TIMEOUT")
        tc.steps.append(Step(
            no=int(g("STEP") or 0),
            type=typ,
            value=str(val),
            method=str(g("TEST_METHOD") or ""),
            fail=str(g("FAIL")) if g("FAIL") else None,
            timeout=float(to) if to is not None else None,
            timeout_from=str(g("TIMEOUT_FROM") or "ACTION").upper(),
            match=str(g("MATCH") or "CONTAINS").upper(),
            scope=str(g("SCOPE") or "").upper(),
            capture=str(g("CAPTURE")) if g("CAPTURE") else None,
        ))

    for tc in cases.values():
        tc.steps.sort(key=lambda s: s.no)

    return list(cases.values())


def summarize(cases: list[TestCase], product: str = "CRX") -> dict:
    mine = [c for c in cases if c.applies_to(product)]
    runnable = [c for c in mine if not c.blocked_by and c.checks]
    blocked = [c for c in mine if c.blocked_by]
    return {
        "total": len(cases),
        "product": len(mine),
        "runnable": len(runnable),
        "blocked": len(blocked),
        "blocked_reasons": _count([str(c.blocked_by) for c in blocked]),
    }


def _count(items: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i in items:
        out[i] = out.get(i, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: -kv[1]))
